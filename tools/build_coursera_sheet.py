#!/usr/bin/env python3
"""Build coursera-prices.xlsx: one tab per offering type, from the two
harvest JSONs (API price sweeps + page crawls)."""
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
BORDER = Border(bottom=Side(style="thin", color="B0B0B0"))
CTRL = re.compile(r"[\x00-\x1f\x7f]")

h2 = json.load(open("research/coursera-price-harvest2.json"))
h1 = json.load(open("research/coursera-price-harvest.json"))

wb = Workbook()


def scaffold(ws, title, note, headers, widths):
    ws["A1"] = title
    ws["A1"].font = Font(name=ARIAL, size=13, bold=True)
    ws["A2"] = note
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True, color="666666")
    for c, head in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=head)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="333F50")
        cell.alignment = Alignment(vertical="center")
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"


def put_row(ws, r, vals, money_cols=(), link_col=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=CTRL.sub("", v) if isinstance(v, str) else v)
        cell.font = Font(name=ARIAL, size=10)
        cell.border = BORDER
    for c in money_cols:
        ws.cell(row=r, column=c).number_format = "$#,##0.00"
    if link_col:
        u = ws.cell(row=r, column=link_col)
        if u.value:
            u.hyperlink = u.value
            u.font = Font(name=ARIAL, size=9, color="0563C1", underline="single")


def summary(ws, s, col, r0, last, extra=()):
    stats = [("Items with a price", f"=COUNT({col}{r0}:{col}{last})", "0"),
             ("Min", f"=MIN({col}{r0}:{col}{last})", "$#,##0.00"),
             ("Median", f"=MEDIAN({col}{r0}:{col}{last})", "$#,##0.00"),
             ("Max", f"=MAX({col}{r0}:{col}{last})", "$#,##0.00")] + list(extra)
    ws.cell(row=s, column=1, value="Summary").font = Font(name=ARIAL, size=11, bold=True)
    for i, (label, formula, fmt) in enumerate(stats, 1):
        ws.cell(row=s + i, column=1, value=label).font = Font(name=ARIAL, size=10)
        c = ws.cell(row=s + i, column=2, value=formula)
        c.font = Font(name=ARIAL, size=10, bold=True)
        c.number_format = fmt


SRC = ("Source: Coursera productPrices.v3 API and page crawls, Aug 12 2026 "
       "(GitHub Actions, matthewkang858/random-cuisine-picker). USD, US storefront.")

# ---- tab 1: read me --------------------------------------------------------
ws = wb.active
ws.title = "Read me"
lines = [
    ("Coursera pricing — all consumer offerings", 13, True),
    (SRC, 9, False),
    ("", 10, False),
    ("How Coursera pricing works: most catalog content is sold by SUBSCRIPTION, not per item.", 10, True),
    ("• Single courses: one-time verified-certificate purchase (Single courses tab; full catalog).", 10, False),
    ("• Specializations: $49/month (range $39-79) until completion, 7-day free trial; 40% also sell a", 10, False),
    ("   prepaid bundle priced at months x $49 (Specializations tab).", 10, False),
    ("• Professional Certificates: $49/month typical (range $39-99), 7-day free trial; a few sell", 10, False),
    ("   prepaid (Professional Certs tab).", 10, False),
    ("• Coursera Plus: $59/month (7-day trial) or $399/year (14-day money-back); covers 90%+ of the", 10, False),
    ("   catalog incl. most Professional Certificates. Promo live Aug 2026: 40% off first 3 months.", 10, False),
    ("• Graduate Certificates and MasterTrack: one-time university tuition, $2,000-$11,000 typical", 10, False),
    ("   (their tabs; not included in Coursera Plus).", 10, False),
    ("• Financial aid is available for courses, Specializations, and most Professional Certificates.", 10, False),
]
for i, (text, size, bold) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=text)
    c.font = Font(name=ARIAL, size=size, bold=bold)
ws.column_dimensions["A"].width = 110

# ---- tab 2: single courses -------------------------------------------------
ws = wb.create_sheet("Single courses")
courses = sorted(
    ((v.get("name") or v.get("slug") or cid, v.get("slug", ""), v.get("amount"))
     for cid, v in h2["courses"].items()),
    key=lambda x: (-(x[2] if x[2] is not None else -1), x[0]))
scaffold(ws, f"Coursera single courses — verified-certificate price ({len(courses):,} courses)",
         SRC + " Price = one-time certificate purchase for that course; auditing is free. "
         "URLs are plain text to keep this 23k-row tab lean.",
         ["Course", "Certificate price ($)", "URL"], [64, 18, 60])
r0 = 5
for i, (name, slug, amount) in enumerate(courses):
    put_row(ws, r0 + i, [name, amount, f"https://www.coursera.org/learn/{slug}"], money_cols=(2,))
last = r0 + len(courses) - 1
summary(ws, last + 2, "B", r0, last, extra=[
    ("Courses at $49 (standard)", f"=COUNTIF(B{r0}:B{last},49)", "0"),
    ("Courses at $9.99 (Google Cloud tier)", f"=COUNTIF(B{r0}:B{last},9.99)", "0"),
    ("Courses at $79+", f'=COUNTIF(B{r0}:B{last},">=79")', "0")])

# ---- tab 3: specializations ------------------------------------------------
ws = wb.create_sheet("Specializations")
spec_rows = []
for sid, meta in h2["spec_map"].items():
    p = h2["spec_prices"].get(sid, {}).get("amount")
    note = ""
    if p is not None and p % 49 == 0:
        note = f"prepaid = {int(p // 49)} months x $49"
    elif p is None:
        note = "subscription only ($49/mo typical)"
    spec_rows.append((meta["name"] or meta["slug"], p, note,
                      f"https://www.coursera.org/specializations/{meta['slug']}"))
spec_rows.sort(key=lambda x: (-(x[1] if x[1] is not None else -1), x[0]))
scaffold(ws, f"Coursera Specializations ({len(spec_rows):,})",
         SRC + " All are $49/month (range $39-79) until completion after a 7-day free trial; "
         "the prepaid column is the buy-outright bundle where offered.",
         ["Specialization", "Prepaid price ($)", "Pricing note", "URL"], [58, 17, 30, 58])
r0 = 5
for i, row in enumerate(spec_rows):
    put_row(ws, r0 + i, list(row), money_cols=(2,))
last = r0 + len(spec_rows) - 1
summary(ws, last + 2, "B", r0, last, extra=[
    ("Subscription-only (no prepaid)", f'=COUNTIF(C{r0}:C{last},"subscription only*")', "0")])

# ---- tab 4: professional certificates -------------------------------------
ws = wb.create_sheet("Professional Certs")
monthly_by_slug = {}
for r in h1.get("profcerts", []):
    slug = r["url"].rsplit("/", 1)[1]
    if r.get("monthly"):
        monthly_by_slug[slug] = min(int(m) for m in r["monthly"])
pc_rows = []
for sid, meta in h2["pc_map"].items():
    p = h2["spec_prices"].get(sid, {}).get("amount")
    m = monthly_by_slug.get(meta["slug"])
    pc_rows.append((meta["name"] or meta["slug"], m, p,
                    "" if m or p else "$49/mo typical (page shows price at enrollment)",
                    f"https://www.coursera.org/professional-certificates/{meta['slug']}"))
pc_rows.sort(key=lambda x: (-(x[2] if x[2] is not None else -1),
                            -(x[1] if x[1] is not None else -1), x[0]))
scaffold(ws, f"Coursera Professional Certificates ({len(pc_rows)})",
         SRC + " Subscription products: $49/month typical (range $39-99) after a 7-day free "
         "trial; most are included in Coursera Plus. Monthly price shown where the program "
         "page displays it; prepaid where the API prices a buy-outright option.",
         ["Professional Certificate", "Monthly ($)", "Prepaid ($)", "Note", "URL"],
         [56, 12, 12, 34, 58])
r0 = 5
for i, row in enumerate(pc_rows):
    put_row(ws, r0 + i, list(row), money_cols=(2, 3))
last = r0 + len(pc_rows) - 1
summary(ws, last + 2, "C", r0, last)

# ---- tabs 5-6: graduate certificates + mastertrack -------------------------
def tuition_tab(name, key, url_note):
    ws = wb.create_sheet(name)
    rows = []
    for r in h1.get(key, []):
        slug = r["url"].rsplit("/", 1)[1]
        program = slug.replace("-", " ").title()
        tuition = r.get("tuition") or []
        shown = " | ".join(tuition[:2]) if tuition else ""
        usd = None
        joined = " ".join(tuition)
        m = re.search(r"\$\s?([0-9]{1,2},[0-9]{3}|[0-9]{3,5})(?!\d)", joined)
        if m:
            usd = float(m.group(1).replace(",", ""))
        if not shown and r.get("monthly"):
            shown = f"${r['monthly'][0]} per month"
        rows.append((program, usd, shown or "tuition not shown on page", r["url"]))
    rows.sort(key=lambda x: (-(x[1] if x[1] is not None else -1), x[0]))
    scaffold(ws, f"Coursera {name} ({len(rows)} programs)",
             SRC + " " + url_note,
             ["Program", "Tuition USD (parsed)", "Tuition as shown on page", "URL"],
             [46, 18, 52, 56])
    r0 = 5
    for i, row in enumerate(rows):
        put_row(ws, r0 + i, list(row), money_cols=(2,))
    last = r0 + len(rows) - 1
    summary(ws, last + 2, "B", r0, last)

tuition_tab("Graduate Certificates", "gradcerts",
            "One-time university tuition, often credit-bearing; not in Coursera Plus. Some pages "
            "render tuition only client-side — those rows are marked.")
tuition_tab("MasterTrack", "mastertrack",
            "One-time university tuition, credit toward the partner degree; not in Coursera Plus. "
            "Some programs price in local currency (shown verbatim).")

# ---- tab 7: coursera plus --------------------------------------------------
ws = wb.create_sheet("Coursera Plus")
scaffold(ws, "Coursera Plus",
         SRC + " Vendor-confirmed on coursera.org/courseraplus.",
         ["Plan", "Price", "Terms"], [26, 16, 80])
plus_rows = [
    ("Monthly", "$59/month", "7-day free trial; cancel anytime; 90%+ of catalog incl. most Professional Certificates"),
    ("Annual", "$399/year", "14-day money-back guarantee; same coverage; ~44% cheaper than 12 months of monthly"),
    ("Promo (Aug 2026)", "40% off first 3 months", "≈$35.40/month via coursera.org promo pages and student channels (UNiDAYS/Student Beans)"),
    ("Excluded", "—", "Degrees, MasterTrack, Graduate Certificates, and select certificates are not included"),
]
r0 = 5
for i, row in enumerate(plus_rows):
    put_row(ws, r0 + i, list(row))

wb.save("research/coursera-prices.xlsx")
print("wrote research/coursera-prices.xlsx")
for s in wb.sheetnames:
    print(" tab:", s)
