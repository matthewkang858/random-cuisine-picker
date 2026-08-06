#!/usr/bin/env python3
"""Build the interactive HTML dashboard from the analysis workbook.

Reads output/tcg_price_trends.xlsx (per-game monthly median/IQR/basket
sheets + the GMV-weighted market composite), embeds the data as JSON into
scripts/dashboard_template.html, and writes a self-contained
output/tcg_dashboard.html — no network, no external libraries.

Usage:
    python3 scripts/build_dashboard.py [--body-only PATH]

--body-only additionally writes the unwrapped page content (no
<!doctype>/<html>/<head>/<body> skeleton) to PATH, for hosts that wrap it
themselves.
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tcg_config import CATEGORIES, CATEGORY_ORDER, GMV_3MO

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "output" / "tcg_price_trends.xlsx"
TEMPLATE = ROOT / "scripts" / "dashboard_template.html"
OUT = ROOT / "output" / "tcg_dashboard.html"

# The six games broken out in the client exhibit tables; slot order is the
# fixed color assignment (market-importance order from tcg_config).
FOCUS = ["Magic", "Pokemon", "Yu-Gi-Oh", "One Piece", "Lorcana", "Riftbound"]

PAGE_SKELETON = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>TCG Market Explorer</title>
</head>
<body>
{body}
</body>
</html>
"""


def rnd(v, nd=2):
    return None if v is None else round(float(v), nd)


def extract(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)

    months = [r[0] for r in wb["Index (base=100)"].iter_rows(min_row=2, values_only=True)]
    mi = {m: i for i, m in enumerate(months)}

    summary = {}
    for seg, base, n, final, cagr in wb["Summary & baskets"].iter_rows(min_row=2, values_only=True):
        summary[seg] = {"base": base, "n": n, "final": rnd(final, 1), "cagr": rnd(cagr, 4)}

    games = []
    for cid in CATEGORY_ORDER:
        short, full = CATEGORIES[cid]
        cols = {k: [None] * len(months) for k in
                ("s_med", "s_p25", "s_p75", "x_med", "x_p25", "x_p75")}
        base = None
        for r in wb[short].iter_rows(min_row=2, values_only=True):
            m = r[0]
            if base is None:
                base = m
            i = mi[m]
            for j, k in enumerate(("s_med", "s_p25", "s_p75", None, "x_med", "x_p25", "x_p75")):
                if k:
                    cols[k][i] = rnd(r[j + 1])
        s_sum = summary[f"{short} singles"]
        x_sum = summary[f"{short} sealed"]
        games.append({
            "key": short,
            "name": full,
            "gmv": GMV_3MO[cid],
            "base": base,
            "focus": short in FOCUS,
            "s": {"med": cols["s_med"], "p25": cols["s_p25"], "p75": cols["s_p75"],
                  "n": s_sum["n"], "final": s_sum["final"], "cagr": s_sum["cagr"]},
            "x": {"med": cols["x_med"], "p25": cols["x_p25"], "p75": cols["x_p75"],
                  "n": x_sum["n"], "final": x_sum["final"], "cagr": x_sum["cagr"]},
        })

    comp = {"market": [], "singles": [], "sealed": [], "equal": []}
    for m, mkt, sng, sld, eq in wb["Market composite"].iter_rows(min_row=2, values_only=True):
        comp["market"].append(rnd(mkt))
        comp["singles"].append(rnd(sng))
        comp["sealed"].append(rnd(sld))
        comp["equal"].append(rnd(eq))

    return {
        "months": months,
        "focus": FOCUS,
        "games": games,
        "composite": comp,
        "note": "30 monthly TCGCSV snapshots, 2024-02-08 through 2026-07-01",
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--body-only", type=Path, default=None,
                    help="also write the unwrapped page content to this path")
    args = ap.parse_args()

    data = extract(XLSX)
    payload = json.dumps(data, separators=(",", ":")).replace("</", "<\\/")

    body = TEMPLATE.read_text()
    marker = "/*__TCG_DATA__*/null"
    if marker not in body:
        sys.exit(f"marker {marker!r} not found in {TEMPLATE}")
    body = body.replace(marker, payload)

    OUT.write_text(PAGE_SKELETON.format(body=body))
    print(f"wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    if args.body_only:
        args.body_only.parent.mkdir(parents=True, exist_ok=True)
        args.body_only.write_text(body)
        print(f"wrote {args.body_only} ({args.body_only.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
